"""Build docs/WindGuard_TX_Data_Flow.docx: flow diagrams of where each dataset comes from,
where it is calculated and how the results reach the screen (run from the repo root or tools/)."""
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.lines import Line2D
from matplotlib.patches import FancyArrowPatch, FancyBboxPatch

ROOT = Path(__file__).resolve().parent.parent
IMG = ROOT / "docs" / "img"
DOCX = ROOT / "docs" / "WindGuard_TX_Data_Flow.docx"

FONT, MONO = "Segoe UI", "Consolas"
plt.rcParams.update({"font.family": FONT, "font.size": 10.5})

INK, INK2, MUTED, LINE = "#1F2933", "#3E4C59", "#616E7C", "#52606D"
PAL = {  # fill, border
    "user": ("#F1F3F5", "#5F6B76"),
    "soil": ("#F7EFE5", "#A0672D"),
    "wx": ("#E8F1FB", "#2C6CB0"),
    "built": ("#F1ECF8", "#6B4FA0"),
    "calc": ("#E6F4EA", "#2F7D4F"),
    "comb": ("#FFF2E2", "#C26A12"),
    "out": ("#EEF0F2", "#323F4B"),
    "warn": ("#FDECEC", "#B83A3A"),
    "panel": ("#FAFBFC", "#CBD2D9"),
}
DASH = (0, (4, 2.5))


class Box:
    def __init__(self, x, y, w, h):
        self.x, self.y, self.w, self.h = x, y, w, h

    @property
    def top(self): return self.y + self.h

    @property
    def right(self): return self.x + self.w

    @property
    def cx(self): return self.x + self.w / 2

    @property
    def cy(self): return self.y + self.h / 2


class Sheet:
    """A matplotlib canvas measured in inches (1 data unit = 1 inch)."""

    def __init__(self, w, h):
        self.w, self.h = w, h
        self.fig = plt.figure(figsize=(w, h), dpi=220)
        self.ax = self.fig.add_axes([0, 0, 1, 1])
        self.ax.set_xlim(0, w)
        self.ax.set_ylim(0, h)
        self.ax.axis("off")
        self.fits = []  # (text, right limit) checked before saving

    def text(self, x, y, s, fs=10.5, color=INK2, weight="normal", ha="left", va="top", family=None,
             style="normal", limit=None, bg=None, z=4):
        t = self.ax.text(x, y, s, fontsize=fs, color=color, fontweight=weight, ha=ha, va=va, zorder=z,
                         family=family or FONT, style=style,
                         bbox=dict(fc=bg, ec="none", pad=1.2) if bg else None)
        if limit is not None:
            self.fits.append((t, limit))
        return t

    def box(self, x, top, w, h, kind, title, lines=(), module=None, dashed=False, fs=10.5, tfs=12, lh=0.215,
            lw=1.5):
        fc, ec = PAL[kind]
        b = Box(x, top - h, w, h)
        self.ax.add_patch(FancyBboxPatch((b.x, b.y), w, h, boxstyle="round,pad=0,rounding_size=0.09",
                                         fc=fc, ec=ec, lw=lw, ls=DASH if dashed else "-", zorder=2))
        tx, ty = x + 0.14, top - 0.11
        self.text(tx, ty, title, fs=tfs, color=INK, weight="semibold", limit=b.right - 0.08)
        ty -= tfs / 72 * 1.38
        if module:
            self.text(tx, ty + 0.01, module, fs=fs - 1.2, color=ec, family=MONO, limit=b.right - 0.08)
            ty -= (fs - 1) / 72 * 1.45
        for ln in lines:
            self.text(tx, ty, ln, fs=fs, limit=b.right - 0.08)
            ty -= lh
        if ty < b.y - 0.02:
            print(f"  ! text runs past the bottom of box '{title}' by {b.y - ty:.2f} in")
        return b

    def panel(self, x, top, w, h, title=None):
        fc, ec = PAL["panel"]
        self.ax.add_patch(FancyBboxPatch((x, top - h), w, h, boxstyle="round,pad=0,rounding_size=0.12",
                                         fc=fc, ec=ec, lw=1, zorder=1))
        if title:
            self.text(x + 0.12, top - 0.1, title, fs=10, color=MUTED, weight="semibold")

    def path(self, pts, color=LINE, dashed=False, lw=1.4, head=True):
        ls = DASH if dashed else "-"
        if len(pts) > 2:
            xs, ys = zip(*pts[:-1])
            self.ax.add_line(Line2D(xs, ys, color=color, lw=lw, ls=ls, zorder=3, solid_capstyle="butt"))
        self.ax.add_patch(FancyArrowPatch(pts[-2], pts[-1], arrowstyle="-|>" if head else "-", mutation_scale=13,
                                          color=color, lw=lw, ls=ls, shrinkA=0, shrinkB=0, zorder=3))

    def label(self, x, y, s, ha="center", va="bottom", color=MUTED, fs=9, bg="white"):
        return self.text(x, y, s, fs=fs, color=color, ha=ha, va=va, style="italic", bg=bg, z=5)

    def legend_swatch(self, x, y, kind, s, dashed=False):
        fc, ec = PAL[kind]
        self.ax.add_patch(FancyBboxPatch((x, y - 0.09), 0.3, 0.2, boxstyle="round,pad=0,rounding_size=0.04",
                                         fc=fc, ec=ec, lw=1.3, ls=DASH if dashed else "-", zorder=2))
        t = self.text(x + 0.38, y + 0.01, s, fs=9.5, color=INK2, va="center")
        self.fig.canvas.draw()
        return x + 0.38 + t.get_window_extent().width / self.fig.dpi + 0.32

    def save(self, name):
        self.fig.canvas.draw()
        for t, limit in self.fits:
            right = self.ax.transData.inverted().transform(t.get_window_extent().get_points())[1][0]
            if right > limit:
                print(f"  ! '{t.get_text()}' overflows its box by {right - limit:.2f} in")
        IMG.mkdir(parents=True, exist_ok=True)
        out = IMG / name
        self.fig.savefig(out, dpi=220, facecolor="white")
        plt.close(self.fig)
        print("wrote", out.relative_to(ROOT))
        return out


# --------------------------------------------------------------------------------------------
# Figure 1: the whole app, left to right
# --------------------------------------------------------------------------------------------
def fig_overview():
    W, H = 15.1, 11.3
    s = Sheet(W, H)
    SX, SW = 0.45, 3.5
    EX, EW = SX + SW + 0.65, 3.2
    KX, KW = EX + EW + 0.85, 2.7
    RX, RW = KX + KW + 0.85, 2.75

    for x, w, t in [(SX, SW, "1   WHERE THE DATA COMES FROM"), (EX, EW, "2   CALCULATED IN YOUR BROWSER"),
                    (KX, KW, "3   COMBINED"), (RX, RW, "4   WHAT YOU SEE")]:
        s.text(x, 11.08, t, fs=11, color=INK, weight="bold")
        s.ax.add_line(Line2D([x, x + w], [10.8, 10.8], color="#9AA5B1", lw=1.2))

    # top row: the point and the Texas check
    U = s.box(SX, 10.62, SW, 0.9, "user", "You pick a point",
              ["map tap · place search (OpenStreetMap)", "GPS · typed latitude, longitude"])
    L = s.box(EX, 10.62, EW, 0.9, "calc", "Location check", ["inside the Texas outline?"], module="app.js · inTexas")
    M = s.box(RX, 10.62, RW, 0.9, "warn", "Refused", ["“covers Texas only” message;", "nothing is downloaded"])
    s.path([(U.right, U.cy), (L.x, U.cy)])
    s.path([(L.right, L.cy), (M.x, L.cy)], color=PAL["warn"][1], dashed=True)
    s.label((L.right + M.x) / 2, L.cy + 0.05, "no: outside Texas", color=PAL["warn"][1])

    bus_y, chan_y, top0, pitch, bh = 9.5, 9.3, 9.13, 1.04, 0.9
    bx = EX + 0.7
    s.path([(bx, L.y), (bx, bus_y)], head=False)
    s.label(bx + 0.08, (L.y + bus_y) / 2 - 0.01, "yes", ha="left", va="center")
    rows = [top0 - i * pitch for i in range(8)]
    mid = [r - bh / 2 for r in rows]
    s.path([(bx, bus_y), (0.2, bus_y), (0.2, mid[7])], head=False)
    s.label((0.2 + EX) / 2 + 0.2, bus_y + 0.03, "your point (lat, lon) goes to every service at once")

    src = [
        ("soil", "USDA-NRCS SSURGO soil survey", ["Soil Data Access web service", "map unit, horizons, WEG, I, T, factor C"], False),
        ("soil", "USGS 3DEP elevation", ["Elevation Point Query Service", "5 points → slope and aspect"], False),
        ("soil", "Land cover", ["NLCD 2021 (MRLC); when blocked,", "Sentinel-2 10 m (Esri / Impact Obs.)"], False),
        ("wx", "gridMET daily weather", ["Univ. of Idaho THREDDS server, 4 km", "rain, grass reference ET₀, Tmax/Tmin"], False),
        ("wx", "ECMWF hourly wind (default)", ["Open-Meteo archive: re-hosted copy", "10 m speed, direction, daily gust"], True),
        ("wx", "NASA MERRA-2 hourly wind", ["NASA POWER API, ~50 km cell", "if chosen, or ECMWF fails or > 20 s"], False),
        ("wx", "NOAA airport records (GHCNh)", ["NCEI data service; IEM copy as backup", "hourly wind and gusts at the airport"], False),
        ("built", "Airport station list (built in)", ["417 stations (NOAA ISD + GHCNh lists)", "6 nearest within 150 km, ASOS first"], False),
    ]
    S = []
    for i, (kind, title, lines, dashed) in enumerate(src):
        S.append(s.box(SX, rows[i], SW, bh, kind, title, lines, dashed=dashed))
        s.path([(0.2, mid[i]), (SX, mid[i])], lw=1.2)

    E1 = s.box(EX, rows[0], EW, bh + pitch, "calc", "Soil engine",
               ["dominant component, surface layer", "texture class → WEG and index I", "Saxton & Rawls water retention",
                "RWEQ EF and SCF (cross-check)", "profile parameters, .SOL export"], module="soil.js")
    E2 = s.box(EX, rows[3], EW, bh + 4 * pitch, "calc", "Wind & climate engine",
               ["nearest airport, quality-checked", "calibrate model to the airport", "(monthly quantile mapping)",
                "erosive hours: wind ≥ 8 m/s", "erosive energy Σ u²(u − 8)", "dryness = (ET₀ − P) / ET₀",
                "monthly risk = energy × dryness", "red month: ≥ 10 % of annual risk", "wind rose, direction, drift",
                "trend test, FAO index, gust days", "", "no airport with ≥ 60 % of hours:", "   model only, uncalibrated"],
               module="wind.js")
    K1 = s.box(KX, rows[0], KW, bh + pitch, "comb", "Erosion-risk model",
               ["E = I × C / 100   (t/ac/yr)", "EI = E / T", "C from SSURGO; else estimated", "class: Low … Extreme"],
               module="app.js · riskModel")
    K2 = s.box(KX, rows[2], KW, bh + 2 * pitch, "comb", "Practice planner",
               ["land use and your answers", "region (MLRA), soil group, risk", "windy months and direction",
                "→ NRCS practices with specs,", "   timing and species"], module="practices.js")
    R_soil = s.box(RX, rows[0], RW, bh, "out", "Soil tab", ["editable surface layer, profile", "CSV and crop-model .SOL export"], lw=2)
    R_over = s.box(RX, rows[1], RW, bh + pitch, "out", "Overview tab",
                   ["risk class, E and EI", "erosive hours, red months,", "prevailing direction",
                    "(wind values from wind engine)"], lw=2)
    R_prac = s.box(RX, rows[3], RW, bh + pitch, "out", "Practices tab",
                   ["NRCS practice codes and specs", "timing calendar", "cover-crop and native species"], lw=2)
    R_wind = s.box(RX, rows[5], RW, bh + pitch, "out", "Wind tab",
                   ["seasonal hours, red months", "wind rose, moisture, trend", "airport vs model comparison"], lw=2)
    s.box(RX, rows[7], RW, bh, "out", "Methods sheet", ["sources used at this point,", "every equation and reference"], lw=2)

    for i in (0, 1):
        s.path([(S[i].right, mid[i]), (EX, mid[i])])
    s.path([(S[2].right, mid[2]), (KX, mid[2])])
    s.label(EX + EW / 2, mid[2] + 0.05, "land-cover class → default land use")
    for i in range(3, 8):
        s.path([(S[i].right, mid[i]), (EX, mid[i])])

    y = (mid[0] + mid[1]) / 2
    s.path([(E1.right, y), (KX, y)])
    s.label((E1.right + KX) / 2, y + 0.05, "I, WEG, T, C")
    cx = E1.right - 0.45
    s.path([(cx, E1.top), (cx, chan_y), (R_soil.cx, chan_y), (R_soil.cx, R_soil.top)])
    s.label(KX + KW / 2, chan_y + 0.03, "soil profile and derived properties")
    s.path([(K1.cx, K1.y), (K2.cx, K2.top)])
    y = (mid[3] + mid[4]) / 2
    s.path([(E2.right, y), (KX, y)])
    s.label((E2.right + KX) / 2, y + 0.05, "windy months")
    s.path([(K1.right, mid[1]), (RX, mid[1])])
    s.label((K1.right + RX) / 2, mid[1] + 0.05, "E, EI, class")
    s.path([(K2.right, y), (RX, y)])
    s.label((K2.right + RX) / 2, y + 0.05, "plan")
    y = (mid[5] + mid[6]) / 2
    s.path([(E2.right, y), (RX, y)])
    s.label(KX + KW / 2, y + 0.05, "hourly and monthly statistics")

    x = SX
    for kind, t in [("user", "You"), ("soil", "Soil and land data"), ("wx", "Weather and wind data"),
                    ("built", "Built into the app"), ("calc", "Calculated in your browser"), ("comb", "Combined"),
                    ("out", "On screen"), ("warn", "Refused")]:
        x = s.legend_swatch(x, 0.62, kind, t)
    x = s.legend_swatch(SX, 0.24, "panel", "solid border: read directly from the agency that produces it")
    x = s.legend_swatch(x, 0.24, "panel", "dashed border: re-hosted copy", dashed=True)
    s.text(x, 0.25, "Statewide map layers are pre-built offline (tools/build_layers.py) and only displayed.",
           fs=9.5, color=MUTED, va="center")
    return s.save("flow_overview.png")


# --------------------------------------------------------------------------------------------
# Figure 2: wind and climate pipeline, with every fallback
# --------------------------------------------------------------------------------------------
def fig_wind():
    W, H = 8.6, 10.75
    s = Sheet(W, H)
    kw = dict(fs=10, tfs=11, lh=0.2)
    LW, xs = 2.7, [0.1, 2.95, 5.8]
    bw = LW - 0.24
    ptop = 10.62

    def chain(x0, items, top):
        """Stacked boxes inside a lane; items = (kind, title, lines, h, dashed, arrow-to-next label, dashed arrow)."""
        boxes, y = [], top
        for kind, title, lines, h, dashed, note, alt in items:
            b = s.box(x0 + 0.12, y, bw, h, kind, title, lines, dashed=dashed, **kw)
            if boxes:
                s.path([(b.cx, boxes[-1][0].y), (b.cx, b.top)], dashed=boxes[-1][1][1])
                if boxes[-1][1][0]:
                    s.label(b.cx, (boxes[-1][0].y + b.top) / 2, boxes[-1][1][0], va="center", fs=8.5)
            boxes.append((b, (note, alt)))
            y = b.y - 0.42
        return [b for b, _ in boxes]

    for x, t, h in [(xs[0], "A   Gridded hourly wind", 3.15), (xs[1], "B   Airport wind, for calibration", 5.65),
                    (xs[2], "C   Daily moisture", 3.95)]:
        s.panel(x, ptop, LW, h, t)
    A = chain(xs[0], [
        ("wx", "ECMWF via Open-Meteo", ["default · 9 km; ERA5 before 2017", "hourly speed, direction; daily gust",
                                       "usually answers in 1–3 s"], 1.05, True, "error, or no answer in 20 s", True),
        ("wx", "NASA MERRA-2", ["NASA POWER API · direct", "0.5° × 0.625° · 2001 on",
                              "no gusts: daily max wind used"], 1.05, False, None, False),
    ], ptop - 0.42)
    B = chain(xs[1], [
        ("built", "Pick the station", ["6 nearest within 150 km", "record covers the period", "ASOS ranked ahead of AWOS"],
         1.0, False, None, False),
        ("wx", "NOAA NCEI GHCNh", ["routine (FM-15) reports", "drop QC flags 2, 3, 6, 7", "wind ≤ 40, gust ≤ 55 m/s",
                                  "nearest report to each hour"], 1.22, False, "coverage < 60 %", True),
        ("wx", "IEM ASOS copy", ["still < 60 %: next station"], 0.68, True, None, False),
        ("wx", "Model at the station's cell", ["same model, same years", "(Open-Meteo or NASA POWER)"], 0.88, False, None, False),
    ], ptop - 0.42)
    C = chain(xs[2], [
        ("wx", "gridMET (Univ. of Idaho)", ["direct · 4 km · daily", "rain, grass ET₀ (ASCE PM),", "Tmax and Tmin"],
         1.05, False, "fails, or < 80 % of days", True),
        ("wx", "Open-Meteo daily (ECMWF)", ["rain, FAO-56 ET₀, mean temp."], 0.68, True, "fails", True),
        ("wx", "NASA POWER daily", ["rain; ET₀ by Hargreaves"], 0.68, False, None, False),
    ], ptop - 0.42)

    q_top = 4.55
    Q = s.box(0.1, q_top, 5.55, 1.27, "calc", "Calibrate to the airport (monthly quantile mapping)",
              ["for each month, rank model-at-station and station hours (101 quantiles)",
               "map each model hour at your point to the station value of equal rank",
               "gusts the same way; no usable airport → model only, uncalibrated"],
              module="wind.js · trainMappers, calibrate", **kw)
    pa, pb, pc = ptop - 3.15, ptop - 5.65, ptop - 3.95
    s.path([(xs[0] + LW / 2, pa), (xs[0] + LW / 2, Q.top)])
    s.label(xs[0] + LW / 2, (pa + Q.top) / 2, "first source that answers:\nhourly wind at your point", va="center", fs=8.5)
    s.path([(xs[1] + LW / 2, pb), (xs[1] + LW / 2, Q.top)])
    s.label(xs[1] + LW / 2 + 0.06, (pb + Q.top) / 2, "station + model at station", va="center", fs=8.5)

    an_top = 2.85
    s.path([(xs[2] + LW / 2, pc), (xs[2] + LW / 2, an_top)])
    s.label(xs[2] + LW / 2, (pc + an_top) / 2 + 0.6, "first source that answers:\ndaily rain, ET₀, temperature",
            va="center", fs=8.5)
    s.path([(Q.cx, Q.y), (Q.cx, an_top)])
    s.label(Q.cx, (Q.y + an_top) / 2, "calibrated hourly wind", va="center", fs=8.5)
    An = s.box(0.1, an_top, 8.4, 1.72, "calc", "Monthly analysis", [], module="wind.js · analyze", **kw)
    left = ["erosive hours: wind ≥ 8 m/s (adjustable)", "erosive energy: Σ u²(u − 8) per month",
            "dryness D = (ET₀ − P) / ET₀, from 0 to 1", "risk index = energy × D; share of the year",
            "red ≥ 10 % of the year, amber ≥ 6.7 %"]
    right = ["erosive wind rose, 16 sectors", "prevailing and resultant transport direction",
             "Fryberger drift potential DP and RDP", "FAO climatic index C",
             "linear trend with significance test", "gust days: daily gust ≥ 17.9 m/s"]
    for col, x in [(left, 0.24), (right, 4.4)]:
        for i, t in enumerate(col):
            s.text(x, an_top - 0.53 - i * 0.2, t, fs=10, limit=(4.3 if x < 1 else 8.42))

    ow, gap = 2.7, (8.4 - 3 * 2.7) / 2
    outs = [("out", "Wind tab", ["charts, rose, moisture, trend"]),
            ("out", "Overview tab", ["erosive hours, red months"]),
            ("comb", "Practice planner", ["calendar against windy months"])]
    for i, (kind, t, lines) in enumerate(outs):
        x = 0.1 + i * (ow + gap)
        b = s.box(x, 0.78, ow, 0.68, kind, t, lines, lw=2 if kind == "out" else 1.5, **kw)
        s.path([(b.cx, An.y), (b.cx, b.top)])
    return s.save("flow_wind.png")


# --------------------------------------------------------------------------------------------
# Figure 3: soil survey to erosion risk
# --------------------------------------------------------------------------------------------
def fig_soil():
    W, H = 8.6, 10.35
    s = Sheet(W, H)
    kw = dict(fs=10, tfs=11, lh=0.2)
    MX, MW = 0.1, 4.5
    SX, SW, chan = 4.95, 3.3, 8.45

    P1 = s.box(MX, 10.25, MW, 0.5, "user", "Your point (inside Texas)", **kw)
    P2 = s.box(MX, 9.4, MW, 0.95, "soil", "SSURGO: soil map unit at the point",
               ["USDA-NRCS Soil Data Access, read directly"], module="SDA_Get_Mukey_from_intersection_with_WktWgs84", **kw)
    P3 = s.box(MX, 8.05, MW, 1.15, "soil", "8 queries at once",
               ["map unit, components, horizons, special areas", "(climatic factor C, rainfall R, MLRA),",
                "ecological site, crop yields, plants, restrictions"], **kw)
    P4 = s.box(MX, 6.5, MW, 0.82, "calc", "Dominant component → surface layer",
               ["sand, silt, clay, OM, CaCO₃, rock fragments"], module="soil.js", **kw)
    P5 = s.box(MX, 5.3, MW, 1.12, "calc", "Texture class → WEG and index I",
               ["USDA texture triangle; NRCS WEG rules: texture,", "clay, CaCO₃ ≥ 5 %, OM ≥ 20 %, stones",
                "I from the WEG table (t/ac/yr)"], module="soil.js · wegFromSurface", **kw)
    P6 = s.box(MX, 3.8, MW, 1.12, "comb", "Bare-soil erosion potential",
               ["E = I × C / 100   (t/ac/yr)", "C: SSURGO climatic factor; if missing,",
                "0.9 × FAO index C from the wind engine"], module="app.js · riskModel", **kw)
    P7 = s.box(MX, 2.3, MW, 1.0, "comb", "Erodibility index and risk class",
               ["EI = E / T   (T = soil-loss tolerance, SSURGO)", "Low < 10 · Moderate 10–40 · High 40–80",
                "Very high 80–150 · Extreme ≥ 150 t/ac/yr"], **kw)
    P8 = s.box(MX, 0.88, MW, 0.78, "comb", "Practice planner", ["+ land use, your answers, region, windy months"],
               module="practices.js", **kw)
    chain = [P1, P2, P3, P4, P5, P6, P7, P8]
    notes = [None, "map unit key", "components and horizons", "surface layer", "I, WEG", "E", "risk, WEG"]
    for a, b, n in zip(chain, chain[1:], notes):
        s.path([(a.cx - 1.2, a.y), (a.cx - 1.2, b.top)])
        if n:
            s.label(a.cx - 1.1, (a.y + b.top) / 2, n, ha="left", va="center", fs=8.5)

    Q1 = s.box(SX, 9.4, SW, 0.95, "soil", "USGS 3DEP elevation", ["5 points: center + 4 at ~100 m", "→ slope %, aspect"], **kw)
    Q2 = s.box(SX, 8.05, SW, 1.15, "calc", "Profile parameters",
               ["all horizons: LL, DUL, SAT, AWC, Ksat,", "root growth, runoff CN, SLU1, SLDR"],
               module="soil.js · toSOL", **kw)
    Q3 = s.box(SX, 6.5, SW, 0.82, "user", "Your edits (Soil tab)", ["sand, silt, clay, OM, CaCO₃, pH"], **kw)
    Q4 = s.box(SX, 5.3, SW, 1.12, "calc", "Edited-layer properties",
               ["Saxton & Rawls 2006 water retention", "RWEQ EF and SCF (cross-check)", "flags soils outside RWEQ range"],
               module="soil.js · saxtonRawls, rweq", **kw)
    Q5 = s.box(SX, 3.8, SW, 1.12, "out", "Soil tab", ["surface layer and full profile", "CSV and crop-model .SOL export"], lw=2, **kw)
    Q6 = s.box(SX, 2.3, SW, 1.0, "out", "Overview tab", ["risk card and 12-month strip", "WEG, C, EI"], lw=2, **kw)
    Q7 = s.box(SX, 0.88, SW, 0.78, "out", "Practices tab", ["NRCS practices, specs, timing"], lw=2, **kw)

    s.path([(Q1.cx, Q1.y), (Q1.cx, Q2.top)])
    s.path([(P3.right, P3.cy), (Q2.x, P3.cy)])
    s.label((P3.right + Q2.x) / 2, P3.cy + 0.05, "horizons", fs=8.5)
    s.path([(Q3.x, Q3.cy), (P4.right, Q3.cy)])
    s.label((P4.right + Q3.x) / 2, Q3.cy + 0.05, "replaces", fs=8.5)
    s.path([(P5.right, P5.cy), (Q4.x, P5.cy)])
    s.label((P5.right + Q4.x) / 2, P5.cy + 0.05, "texture", fs=8.5)
    s.path([(Q4.cx, Q4.y), (Q4.cx, Q5.top)])
    s.path([(Q2.right, Q2.cy), (chan, Q2.cy), (chan, Q5.cy), (Q5.right, Q5.cy)])
    s.path([(P7.right, P7.cy), (Q6.x, P7.cy)])
    s.path([(P8.right, P8.cy), (Q7.x, P8.cy)])
    return s.save("flow_soil.png")


# --------------------------------------------------------------------------------------------
# The Word document
# --------------------------------------------------------------------------------------------
SOURCES = [
    ("Soil survey: map unit, components, horizons, WEG, I, T, climatic factor C, rainfall R, MLRA, yields, plants",
     "USDA-NRCS SSURGO", "Soil Data Access SQL web service (direct)", "survey polygons, mapped at 1:12,000–1:24,000",
     "Soil tab shows the error; risk cannot be computed"),
    ("Elevation, slope and aspect", "USGS 3DEP", "Elevation Point Query Service, 5 points (direct)",
     "best available 3DEP: 1 m lidar where flown, else ~10 m", "Copernicus 90 m elevation via Open-Meteo (after 15 s)"),
    ("Land cover", "MRLC NLCD 2021; Esri / Impact Observatory", "WMS GetFeatureInfo; ImageServer identify (direct)",
     "30 m; Sentinel-2 10 m", "Sentinel-2 when NLCD is blocked (currently always: MRLC refuses web-page requests)"),
    ("Daily rain, grass reference ET₀, Tmax, Tmin", "University of Idaho gridMET", "THREDDS OPeNDAP, one cell (direct)",
     "4 km, daily, 1979 on; last ~4 days not yet published", "Open-Meteo daily (ECMWF), then NASA POWER daily with Hargreaves ET₀"),
    ("Hourly 10 m wind speed and direction, daily max gust (default)", "ECMWF: IFS analysis from 2017, ERA5 before",
     "Open-Meteo archive (re-hosted copy); two requests in parallel", "9 km from 2017; 0.25° (~28 km) before",
     "NASA MERRA-2 if Open-Meteo errors or has not answered within 20 s"),
    ("Hourly 10 m wind speed and direction (alternative)", "NASA GMAO MERRA-2", "NASA POWER hourly API (direct)",
     "0.5° × 0.625° (~50 km), 2001 on; no gust variable", "ECMWF via Open-Meteo"),
    ("Airport hourly wind and gusts (for calibration)", "NOAA NCEI GHCNh", "NCEI Access Data Service, one request per year, 6 at a time (direct)",
     "station, hourly", "Iowa Environmental Mesonet ASOS copy; then the next station"),
    ("Airport station list", "NOAA ISD history + GHCNh station list", "built into the app (data/stations.js)",
     "417 Texas-area stations", "—"),
    ("Place search", "OpenStreetMap Nominatim", "search API; results outside Texas are dropped", "—", "type coordinates instead"),
    ("Statewide map layers", "SSURGO WEG/I (UC Davis SoilWeb), NRCS C, NASA POWER wind; built offline",
     "built into the app (data/layers.js and PNG overlays)", "statewide grid", "display only"),
    ("Base maps", "Esri, USGS, OpenStreetMap", "map tiles", "—", "switch base map"),
]

WIND_CALC = [
    ("Pick the station", "6 nearest within 150 km whose record covers the period; AWOS ranked as if 100 km farther",
     "station list → candidates", "wind.js · `nearestStations`"),
    ("Clean the airport record", "routine FM-15 reports only; drop quality flags 2, 3, 6, 7; wind ≤ 40 m/s, gust ≤ 55 m/s; "
     "keep the report nearest each whole hour; need ≥ 60 % of hours", "NOAA reports → hourly series, coverage",
     "wind.js · `toHourly`, `ghcnhObs`"),
    ("Calibrate the model", "monthly empirical quantile mapping, 101 quantiles from 0.1 to 99.9 %; beyond the ends, scale "
     "by the end ratio; gusts the same way (MERRA-2: daily max wind)", "model at station + station, same years → mapping "
     "applied at your point", "wind.js · `trainMappers`, `calibrate`"),
    ("Erosive hours", "hours with u ≥ uₜ (default 8 m/s, adjustable), scaled to a full month", "hourly wind → hours/month",
     "wind.js · `analyze`"),
    ("Erosive energy", "Σ u² (u − uₜ) over erosive hours", "hourly wind → energy/month", "`analyze`"),
    ("Dryness", "D = (ET₀ − P) / ET₀, limited to 0–1", "monthly rain, ET₀ → D", "`analyze`"),
    ("Monthly risk and red months", "index = energy × D; share = month ÷ year; red ≥ 10 %, amber ≥ 6.7 %",
     "energy, D → 12-month strip", "`analyze`"),
    ("Wind rose and directions", "16 sectors; prevailing = sector with most erosive energy (neighbours weighted ½); "
     "resultant = vector sum of energy toward downwind", "speed, direction → rose, directions", "`analyze`"),
    ("Drift potential", "Fryberger: Σ V² (V − 12 kn) for V > 12 kn; RDP = vector sum; RDP/DP", "hourly wind → DP, RDP",
     "`analyze`"),
    ("FAO climatic index", "C = Σ u₂³ · D · days / 100, with u₂ = 0.748 u₁₀ (FAO-56 log law)", "monthly wind, D → C",
     "`analyze`"),
    ("Gust days", "days with daily max gust ≥ 17.9 m/s (40 mph)", "daily gust → days/month", "`analyze`"),
    ("Trend", "least-squares slope of yearly erosive hours and mean wind, with significance test", "yearly totals → trend",
     "util.js · `linreg`"),
    ("Reference ET (last fallback)", "Hargreaves: ET₀ = 0.0023 Ra (Tmean + 17.8) √(Tmax − Tmin)", "NASA POWER daily → ET₀",
     "wind.js · `powerDaily`"),
]

SOIL_CALC = [
    ("Find the map unit", "`SDA_Get_Mukey_from_intersection_with_WktWgs84` at the point", "lat, lon → map unit key",
     "soil.js · `fetchSoil`"),
    ("Dominant component, surface layer", "component with the largest percentage; its top horizon", "components, horizons → "
     "sand, silt, clay, OM, CaCO₃", "soil.js"),
    ("Texture class", "USDA texture triangle, refined by sand size", "sand, silt, clay → class",
     "soil.js · `textureClass`, `refineCode`"),
    ("Wind Erodibility Group and I", "NRCS rules: texture family, clay ≤ 5 %, CaCO₃ ≥ 5 %, OM ≥ 20 %, stones; I from the WEG table",
     "surface layer → WEG 1–8, I (t/ac/yr)", "soil.js · `wegFromSurface`"),
    ("Climatic factor C", "SSURGO special-area overlap “climatic factor”; if missing, 0.9 × FAO index C (≤ 200)",
     "SSURGO or wind engine → C", "app.js · `cFactor`"),
    ("Bare-soil erosion potential", "E = I × C / 100", "I, C → E (t/ac/yr)", "app.js · `riskModel`"),
    ("Erodibility index", "EI = E / T", "E, T → EI", "`riskModel`"),
    ("Risk class", "Low < 10, Moderate 10–40, High 40–80, Very high 80–150, Extreme ≥ 150 t/ac/yr", "E → class",
     "`riskModel`"),
    ("Water retention", "Saxton & Rawls (2006) pedotransfer equations", "sand, clay, OM → LL, DUL, SAT, AWC, Ksat",
     "soil.js · `saxtonRawls`"),
    ("RWEQ erodible fraction", "EF = (29.09 + 0.31 Sa + 0.17 Si + 0.33 Sa/Cl − 2.59 OM − 0.95 CaCO₃) / 100",
     "surface layer → EF (cross-check)", "soil.js · `rweq`"),
    ("RWEQ crust factor", "SCF = 1 / (1 + 0.0066 Cl² + 0.021 OM²)", "clay, OM → SCF (cross-check)", "soil.js · `rweq`"),
    ("Slope and aspect", "finite differences on the centre and four neighbours ~100 m away", "5 elevations → slope %, aspect",
     "app.js · `fetchElevation`"),
    ("Profile parameters", "SLU1, SLDR, runoff curve number, root growth factor, albedo per crop-model .SOL format",
     "all horizons → .SOL file", "soil.js · `toSOL`"),
]

RESULTS = [
    ("Risk class and E (Overview)", "E = I × C / 100", "SSURGO: WEG → I, climatic factor C"),
    ("Erodibility index EI", "E / T", "SSURGO soil-loss tolerance T"),
    ("Erosive hours, mean wind", "count of hours ≥ threshold; mean", "ECMWF (Open-Meteo) or MERRA-2, calibrated to the NOAA airport"),
    ("Red and amber months", "energy × dryness, share of the year", "calibrated wind + gridMET rain and ET₀"),
    ("Prevailing erosive direction", "sector with most erosive energy", "calibrated wind direction"),
    ("Wind Erodibility Group, C, land cover", "read as published (WEG recomputed if you edit)", "SSURGO; NLCD or Sentinel-2"),
    ("Soil profile, water properties", "Saxton & Rawls for edited layers", "SSURGO horizons"),
    ("Slope", "finite differences", "USGS 3DEP"),
    ("Moisture chart", "monthly sums", "gridMET (or its fallbacks)"),
    ("Airport vs model card", "same-hour comparison", "NOAA GHCNh and the model at the station's cell"),
    ("Practice plan", "rule tables in practices.js", "land use and answers, MLRA region, WEG, risk, windy months"),
    ("Statewide map colours", "pre-computed offline", "SSURGO WEG/I (UC Davis SoilWeb rasters), NRCS C, NASA POWER monthly wind"),
]

CODE = [
    ("index.html, css/app.css", "page layout, tabs, sheets and the phone layout"),
    ("js/util.js", "downloads with time limits (`fetchJSON`, `fetchText`), concurrency pool, statistics helpers"),
    ("js/soil.js", "SSURGO queries, texture, WEG and I, Saxton & Rawls, RWEQ, CSV and .SOL export"),
    ("js/wind.js", "wind and moisture downloads with fallbacks, station choice and QC, calibration, monthly analysis"),
    ("js/app.js", "map, Texas check, loading sequence, risk model, elevation and land cover, rendering of every tab"),
    ("js/charts.js", "Chart.js charts"),
    ("js/practices.js", "practice planner: questions, rules, specifications, species tables"),
    ("js/methods.js", "Methods sheet: sources used at the point, source catalog, equations"),
    ("data/stations.js", "built-in airport list (made by tools/fetch_stations.py)"),
    ("data/layers.js, data/*.png", "statewide layers and the Texas outline (made by tools/build_layers.py)"),
    ("tools/*.py", "offline builders, run by hand, never by the app; tools/build_flow_doc.py makes this document"),
]

STEPS = [
    ("Location check.", " The point is tested against the Texas outline drawn on the map (U.S. Census 1:5,000,000 state "
     "boundary). Outside it, the app shows a message and downloads nothing. Place search lists only places inside Texas."),
    ("Soil.", " Soil Data Access finds the SSURGO map unit under the point, then runs 8 queries at once. The soil engine "
     "derives texture, WEG, I and water properties, and the risk card can show E as soon as the soil arrives."),
    ("Elevation and land cover, at the same time.", " Five USGS 3DEP elevations give slope and aspect (15 s limit, then "
     "Copernicus via Open-Meteo). Land cover comes from NLCD (8 s limit) or the Sentinel-2 10 m map."),
    ("Gridded wind and moisture, at the same time.", " Hourly wind for the chosen period comes from ECMWF through "
     "Open-Meteo (usually 1–3 s) or NASA MERRA-2. If Open-Meteo errors or has not answered within 20 s, the app switches "
     "to MERRA-2. Daily rain and ET₀ come from gridMET. The Wind tab first shows the model, uncalibrated."),
    ("Airport calibration, in the background.", " The app takes up to 6 nearby airports from its built-in list, downloads "
     "hourly records from NOAA one year per request, keeps the first station with at least 60 % of hours, downloads the "
     "same model at that station's cell, trains the monthly quantile mapping and redraws every result as calibrated."),
    ("Practices.", " When you choose a land use (the default comes from land cover) and answer 2–5 questions, the "
     "planner builds the plan from the risk, soil group, region and windy months."),
    ("Changes.", " Editing the soil layer, the threshold or the units recalculates in the browser with no new downloads. "
     "Changing the period, the model or the calibration station downloads again."),
]

LIMITS = [
    ("Texas only.", " The outline is the Census 1:5,000,000 boundary, kept within about 0.2 km of that source; a point a "
     "few hundred metres from the state line can land on the wrong side."),
    ("ECMWF comes through a copy.", " Copernicus's own data store needs each user's personal key and a queued job that "
     "takes minutes, so the app reads Open-Meteo's copy, which arrives in seconds. NASA MERRA-2 is the direct alternative "
     "and the automatic fallback."),
    ("NLCD is blocked.", " MRLC's server refuses requests from web pages, so land cover currently always comes from the "
     "Sentinel-2 10 m map."),
    ("MERRA-2 has no gusts.", " Its daily maximum hourly wind stands in for gusts during calibration, and gust days show n/a."),
    ("gridMET lags a few days.", " The newest ~4 days are not yet published; a point with less than 80 % of days falls "
     "back to Open-Meteo and then NASA POWER."),
    ("Station rules.", " A station needs at least 60 % of hours in the period (20 % if you pick it yourself). AWOS sites are "
     "ranked as if 100 km farther, because several West Texas AWOS anemometers read 40–60 % fewer erosive hours than "
     "nearby ASOS sites."),
    ("RWEQ values are cross-checks.", " EF and SCF are not used in the risk; they get an asterisk when the soil is "
     "outside the range the equations were fitted on."),
    ("Statewide layers are fixed.", " The map colours are pre-computed and do not change with the point you pick."),
    ("No WindGuard server.", " GitHub Pages only delivers the files. Your browser sends the point's coordinates (and any "
     "place name you search for) straight to each service listed above; nothing is stored."),
]


def build_docx(figs, commit):
    from docx import Document
    from docx.enum.section import WD_ORIENT, WD_SECTION
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt, RGBColor

    ink, accent, muted = RGBColor(0x1F, 0x29, 0x33), RGBColor(0x2F, 0x7D, 0x4F), RGBColor(0x61, 0x6E, 0x7C)
    doc = Document()

    st = doc.styles
    st["Normal"].font.name, st["Normal"].font.size = "Calibri", Pt(10.5)
    st["Normal"].paragraph_format.space_after = Pt(5)
    st["Normal"].paragraph_format.line_spacing = 1.1
    for name, size, color, before in [("Title", 24, ink, 0), ("Heading 1", 15, accent, 16), ("Heading 2", 12, ink, 10)]:
        f = st[name].font
        f.name, f.size, f.color.rgb, f.bold = "Calibri", Pt(size), color, True
        rf = st[name].element.get_or_add_rPr().get_or_add_rFonts()  # drop theme fonts so Calibri applies
        for a in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
            rf.attrib.pop(qn(a), None)
        rf.set(qn("w:ascii"), "Calibri")
        rf.set(qn("w:hAnsi"), "Calibri")
        st[name].paragraph_format.space_before = Pt(before)
        st[name].paragraph_format.space_after = Pt(4)
    st["Caption"].font.size, st["Caption"].font.color.rgb, st["Caption"].font.italic = Pt(9), muted, True

    def page(section, landscape=False, margin=0.75):
        w, h = Inches(8.5), Inches(11)
        section.orientation = WD_ORIENT.LANDSCAPE if landscape else WD_ORIENT.PORTRAIT
        section.page_width, section.page_height = (h, w) if landscape else (w, h)
        for side in ("left_margin", "right_margin"):
            setattr(section, side, Inches(margin))
        section.top_margin = section.bottom_margin = Inches(0.45 if landscape else 0.75)
        section.footer_distance = Inches(0.25)

    def runs(p, text, size=None):
        """Plain text with `code` spans in Consolas."""
        for i, part in enumerate(text.split("`")):
            if not part:
                continue
            r = p.add_run(part)
            if i % 2:
                r.font.name = "Consolas"
                r.font.size = Pt((size or 10.5) - 1)
            elif size:
                r.font.size = Pt(size)
        return p

    def para(text, style=None, size=None, color=None, align=None, after=None):
        p = doc.add_paragraph(style=style)
        runs(p, text, size)
        if color is not None:
            for r in p.runs:
                r.font.color.rgb = color
        if align is not None:
            p.alignment = align
        if after is not None:
            p.paragraph_format.space_after = Pt(after)
        return p

    def lead_list(items, style):
        for lead, rest in items:
            p = doc.add_paragraph(style=style)
            p.add_run(lead).bold = True
            runs(p, rest)
            p.paragraph_format.space_after = Pt(3)

    def figure(path, width, caption):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(2)
        p.add_run().add_picture(str(path), width=Inches(width))
        para(caption, style="Caption", align=WD_ALIGN_PARAGRAPH.CENTER, after=8)

    def shade(cell, fill):
        pr = cell._tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        for k, v in (("w:val", "clear"), ("w:color", "auto"), ("w:fill", fill)):
            shd.set(qn(k), v)
        pr.append(shd)

    def table(headers, rows, widths, size=9):
        t = doc.add_table(rows=1, cols=len(headers))
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.autofit = False
        tblPr = t._tbl.tblPr
        borders = OxmlElement("w:tblBorders")
        for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
            e = OxmlElement(f"w:{edge}")
            for k, v in (("w:val", "single"), ("w:sz", "4"), ("w:space", "0"), ("w:color", "C9D1D9")):
                e.set(qn(k), v)
            borders.append(e)
        tblPr.insert_element_before(borders, "w:shd", "w:tblLayout", "w:tblCellMar", "w:tblLook", "w:tblCaption",
                                    "w:tblDescription")  # schema order matters to Word
        grid = t._tbl.tblGrid
        for gc, w in zip(grid.findall(qn("w:gridCol")), widths):
            gc.set(qn("w:w"), str(int(w * 1440)))
        hdr = t.rows[0]
        trPr = hdr._tr.get_or_add_trPr()
        rep = OxmlElement("w:tblHeader")
        rep.set(qn("w:val"), "true")
        trPr.append(rep)
        for c, h, w in zip(hdr.cells, headers, widths):
            c.width = Inches(w)
            shade(c, "323F4B")
            p = c.paragraphs[0]
            r = p.add_run(h)
            r.bold, r.font.size, r.font.color.rgb = True, Pt(size), RGBColor(0xFF, 0xFF, 0xFF)
        for k, row in enumerate(rows):
            tr = t.add_row()
            tr._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))  # keep each row on one page
            cells = tr.cells
            for j, (c, v, w) in enumerate(zip(cells, row, widths)):
                c.width = Inches(w)
                if k % 2:
                    shade(c, "F5F7FA")
                p = c.paragraphs[0]
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.0
                runs(p, v, size)
                if j == 0:
                    for r in p.runs:
                        r.bold = True
        spacer = doc.add_paragraph()
        spacer.paragraph_format.space_after = Pt(2)
        return spacer

    # ---- page 1 ----
    sec = doc.sections[0]
    page(sec)
    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = footer.add_run("WindGuard TX · data flow · page ")
    r.font.size, r.font.color.rgb = Pt(8.5), muted
    # PAGE as a complex field (begin / instruction / separate / result / end); Word stalls exporting a fldSimple here
    for kind, text in (("begin", None), (None, " PAGE "), ("separate", None), (None, "1"), ("end", None)):
        r = footer.add_run()
        r.font.size, r.font.color.rgb = Pt(8.5), muted
        if kind:
            fc = OxmlElement("w:fldChar")
            fc.set(qn("w:fldCharType"), kind)
            r._r.append(fc)
        elif text.strip() == "PAGE":
            it = OxmlElement("w:instrText")
            it.set(qn("xml:space"), "preserve")
            it.text = text
            r._r.append(it)
        else:
            r.text = text

    doc.add_heading("WindGuard TX: data flow", level=0)
    para("Where each dataset comes from, where it is calculated, and how the results reach the screen",
         size=13, color=muted, after=2)
    para(f"App: barsa1ly.github.io/windguard-tx   ·   Code: github.com/barsa1ly/windguard-tx (commit {commit})   ·   "
         f"October 28, 2026", size=9, color=muted, after=10)

    doc.add_heading("1. In one paragraph", level=1)
    para("WindGuard TX is a static web page. When you pick a point inside Texas, your browser asks each agency's web "
         "service for the data at that point, all at the same time. Every calculation then runs in your browser, in the "
         "JavaScript files under js/. There is no WindGuard server: GitHub Pages only delivers the files. The soil "
         "survey gives the erodibility (I) and the climatic factor (C), which combine into the bare-soil erosion "
         "potential E = I × C / 100. Hourly wind, calibrated to the nearest airport and weighted by how dry each month "
         "is, decides which months are risky. The practice planner turns both into a plan.")

    doc.add_heading("2. How to read the diagrams", level=1)
    lead_list([
        ("Columns, left to right:", " where the data comes from, what the browser calculates, what is combined, and "
         "what you see."),
        ("Colours:", " brown = soil and land data, blue = weather and wind data, purple = built into the app, "
         "green = calculated in your browser, orange = combined, dark grey = results on screen, red = refused."),
        ("Borders:", " a solid border means the app reads the data directly from the agency that produces it; a "
         "dashed border means a re-hosted copy (Open-Meteo, Iowa Environmental Mesonet)."),
        ("Arrows:", " solid arrows carry data; dashed arrows are fallbacks, taken only when the source above fails, "
         "with the condition written beside them."),
    ], "List Bullet")

    doc.add_heading("3. What happens when you pick a point", level=1)
    lead_list(STEPS, "List Number")

    # ---- figure 1, landscape ----
    sec = doc.add_section(WD_SECTION.NEW_PAGE)
    page(sec, landscape=True, margin=0.5)
    doc.add_heading("4. The whole picture", level=1).paragraph_format.space_before = Pt(0)
    figure(figs[0], 9.25, "Figure 1. Data sources (left), calculations in the browser (middle) and results on screen "
                         "(right) for one point.")

    # ---- detail pages ----
    sec = doc.add_section(WD_SECTION.NEW_PAGE)
    page(sec)
    doc.add_heading("5. Wind and climate: three sources to monthly risk", level=1).paragraph_format.space_before = Pt(0)
    para("Three lanes run at the same time. Each lane is a priority list: the app uses the first source that answers. "
         "Lane B only exists to calibrate lane A; if no airport qualifies, the model is used uncalibrated.")
    figure(figs[1], 5.9, "Figure 2. Wind and moisture pipeline with every fallback.")

    h = doc.add_heading("6. Soil survey to erosion risk", level=1)
    h.paragraph_format.page_break_before, h.paragraph_format.space_before = True, Pt(0)
    para("The main column is the path to the risk class. The right column holds what the Soil tab adds and the "
         "result boxes. Your edits replace the surveyed surface layer, so every value below them is recalculated.")
    figure(figs[2], 6.2, "Figure 3. Soil and erosion-risk pipeline.")

    h = doc.add_heading("7. Calculations, step by step", level=1)
    h.paragraph_format.page_break_before, h.paragraph_format.space_before = True, Pt(0)
    para("“Where” names the file and function in the app's code, so each step can be checked against the source. "
         "The same equations, with their terms and references, are in the app's Methods sheet and in docs/METHODS.md.")
    doc.add_heading("Wind and climate (Figure 2)", level=2)
    table(["Step", "Rule or equation", "Inputs → output", "Where"], WIND_CALC, [1.35, 3.0, 1.55, 1.1])
    doc.add_heading("Soil and erosion risk (Figure 3)", level=2)
    table(["Step", "Rule or equation", "Inputs → output", "Where"], SOIL_CALC, [1.35, 3.0, 1.55, 1.1])

    doc.add_heading("8. Data sources", level=1)
    table(["Data", "Producer", "How the app reads it", "Resolution", "If it fails"], SOURCES, [1.6, 1.45, 1.45, 1.2, 1.3],
          size=8.5)

    doc.add_heading("9. Where each number on screen comes from", level=1)
    table(["Result on screen", "Calculation", "Data behind it"], RESULTS, [2.0, 2.2, 2.8])

    doc.add_heading("10. Fallbacks, limits and checks", level=1)
    lead_list(LIMITS, "List Bullet")

    doc.add_heading("11. Code map", level=1)
    end = table(["File", "What it does"], CODE, [1.9, 5.1])
    end.paragraph_format.space_after = Pt(0)  # Word needs a paragraph after the last table; keep it 1 pt tall
    mark = OxmlElement("w:rPr")
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), "2")
    mark.append(sz)
    end._p.get_or_add_pPr().append(mark)

    zoom = doc.settings.element.find(qn("w:zoom"))  # python-docx's template omits the required percent
    if zoom is not None and zoom.get(qn("w:percent")) is None:
        zoom.set(qn("w:percent"), "100")
    doc.core_properties.title = "WindGuard TX: data flow"
    doc.core_properties.subject = "Data sources, calculations and results"
    doc.core_properties.author = "WindGuard TX"
    doc.save(DOCX)
    print("wrote", DOCX.relative_to(ROOT))


if __name__ == "__main__":
    import subprocess

    commit = subprocess.run(["git", "-C", str(ROOT), "rev-parse", "--short", "HEAD"], capture_output=True,
                            text=True).stdout.strip() or "unknown"
    build_docx([fig_overview(), fig_wind(), fig_soil()], commit)
